from __future__ import annotations

import csv
import json
import re
from pathlib import Path
import unicodedata
from urllib.parse import urlparse
from urllib.request import Request, urlopen


CENTRAL_MARKETS_PDF_URL = (
    "http://www.protlcuem.gob.mx/work/models/Prologyca/Resource/2/1/images/"
    "DirectorioNacionaldeCentralesdeAbasto.pdf"
)
SCIAN_INDEX_URL = "https://www.inegi.org.mx/scian/"
SCIAN_XLSX_URL = "https://www.inegi.org.mx/contenidos/app/scian/scian_2023_categorias_y_productos.xlsx?v=20230719"

SCIAN_FIELD_ALIASES = {
    "code": "code",
    "codigo": "code",
    "codigo_scian": "code",
    "codigo_scian_2023": "code",
    "clave": "code",
    "title": "title",
    "titulo": "title",
    "categoria": "title",
    "description": "description",
    "descripcion": "description",
    "level": "level",
    "nivel": "level",
    "parent_code": "parent_code",
    "codigo_padre": "parent_code",
    "padre": "parent_code",
}


def load_market_source(path: str) -> list[dict[str, object]]:
    source_path = Path(path)
    suffix = source_path.suffix.lower()

    if suffix == ".pdf":
        return extract_market_rows_from_pdf(source_path)

    rows = _load_structured_rows(source_path)
    normalized_rows: list[dict[str, object]] = []
    for index, row in enumerate(rows, start=1):
        normalized_rows.append(
            {
                "market_id": row.get("market_id") or f"ROW-{index}",
                "market_name": row.get("market_name") or row.get("name") or "",
                "address_raw": row.get("address_raw") or row.get("address") or "",
                "state": row.get("state") or "",
                "municipality": row.get("municipality") or "",
                "source_page": int(row.get("source_page", 0) or 0),
                "source_text": row.get("source_text") or "",
                "lat": row.get("lat"),
                "lon": row.get("lon"),
                "geocode_confidence": row.get("geocode_confidence"),
                "geocode_status": row.get("geocode_status"),
            }
        )

    return normalized_rows


def download_central_markets_pdf(destination_dir: str, *, force: bool = False) -> str:
    directory = Path(destination_dir)
    directory.mkdir(parents=True, exist_ok=True)
    target_path = directory / "DirectorioNacionaldeCentralesdeAbasto.pdf"
    if target_path.exists() and not force:
        return str(target_path)

    _download_to_file(
        url=CENTRAL_MARKETS_PDF_URL,
        target_path=target_path,
        user_agent="ceda-central-markets/0.1",
    )
    return str(target_path)


def download_scian_source(destination_dir: str, *, force: bool = False) -> str:
    directory = Path(destination_dir)
    directory.mkdir(parents=True, exist_ok=True)

    manifest_path = directory / "scian_source_manifest.json"
    if manifest_path.exists() and not force:
        payload = json.loads(manifest_path.read_text(encoding="utf-8"))
        cached_path = directory / str(payload["local_file"])
        if cached_path.exists():
            return str(cached_path)

    source_url = SCIAN_XLSX_URL
    filename = Path(urlparse(source_url).path).name or "scian_source.xlsx"
    downloaded_path = directory / filename
    if not downloaded_path.exists() or force:
        _download_to_file(
            url=source_url,
            target_path=downloaded_path,
            user_agent="ceda-central-markets/0.1",
        )

    manifest_payload = {
        "index_url": SCIAN_INDEX_URL,
        "download_url": source_url,
        "local_file": downloaded_path.name,
    }
    manifest_path.write_text(
        json.dumps(manifest_payload, ensure_ascii=True, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    return str(downloaded_path)


def load_denue_rows(path: str, *, year: int) -> list[dict[str, object]]:
    rows = _load_structured_rows(Path(path))
    filtered: list[dict[str, object]] = []
    for row in rows:
        row_year = row.get("year")
        if row_year in {None, ""}:
            filtered.append(row)
            continue

        if int(row_year) == year:
            filtered.append(row)

    return filtered


def load_scian_rows(path: str) -> list[dict[str, object]]:
    rows = _load_structured_rows(Path(path))

    merged_sectors: list[tuple[int, int, str]] = []
    for row in rows:
        for key, value in row.items():
            if SCIAN_FIELD_ALIASES.get(_normalize_key(key)) != "code":
                continue
            code = _normalize_scian_code(value)
            if code and "-" in code:
                low, high = code.split("-")
                merged_sectors.append((int(low), int(high), code))

    normalized: list[dict[str, object]] = []
    for row in rows:
        renamed: dict[str, object] = {}
        for key, value in row.items():
            alias = SCIAN_FIELD_ALIASES.get(_normalize_key(key))
            if alias:
                renamed[alias] = value

        code = _normalize_scian_code(renamed.get("code"))
        title = _clean_scian_title(str(renamed.get("title", "")))
        if not code or not title:
            continue

        level_raw = renamed.get("level")
        level = int(level_raw) if level_raw not in {None, ""} else len(code)
        parent_raw = _normalize_scian_code(renamed.get("parent_code"))
        parent_code = parent_raw if parent_raw else _infer_parent_code(code, merged_sectors)

        normalized.append(
            {
                "code": code,
                "title": title,
                "description": str(renamed.get("description", "")).strip() or None,
                "level": level,
                "parent_code": parent_code,
            }
        )

    if not normalized:
        raise ValueError("No valid SCIAN rows found in source file")

    return normalized


def extract_market_rows_from_pdf(path: Path) -> list[dict[str, object]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportError("pypdf is required for PDF ingestion. Install dependencies first.") from exc

    reader = PdfReader(str(path))
    results: list[dict[str, object]] = []
    row_index = 0
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        for line in text.splitlines():
            compact = re.sub(r"\s+", " ", line).strip()
            if not compact:
                continue

            parts = _split_market_line(compact)
            if len(parts) < 4:
                continue

            row_index += 1
            state, municipality, market_name = parts[0], parts[1], parts[2]
            address = " ".join(parts[3:])
            results.append(
                {
                    "market_id": f"PDF-{page_number}-{row_index}",
                    "market_name": market_name,
                    "address_raw": address,
                    "state": state,
                    "municipality": municipality,
                    "source_page": page_number,
                    "source_text": compact,
                }
            )

    if not results:
        raise ValueError(
            "No market rows were extracted from PDF. Provide a cleaned CSV/JSON extract or adjust parser heuristics."
        )

    return results


def _split_market_line(line: str) -> list[str]:
    for delimiter in ("|", ";", "\t"):
        if delimiter in line:
            return [piece.strip() for piece in line.split(delimiter) if piece.strip()]

    return [piece.strip() for piece in re.split(r"\s{2,}", line) if piece.strip()]


def _load_structured_rows(path: Path) -> list[dict[str, object]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8", newline="") as handle:
            return list(csv.DictReader(handle))

    if suffix == ".json":
        with path.open("r", encoding="utf-8") as handle:
            payload = json.load(handle)
        if isinstance(payload, list):
            return [dict(item) for item in payload]
        raise ValueError("Expected JSON array of objects")

    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportError("openpyxl is required for XLSX ingestion. Install dependencies first.") from exc

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)

        headers: list[str] | None = None
        parsed_rows: list[dict[str, object]] = []
        for values in rows:
            cells = ["" if value is None else str(value).strip() for value in values]
            if headers is None:
                if any(cells) and _looks_like_scian_header_row(cells):
                    headers = cells
                continue

            if _looks_like_scian_header_row(cells):
                # Handle repeated header rows in long sheets.
                headers = cells
                continue

            if not any(cells):
                continue

            row = {
                headers[index]: values[index]
                for index in range(min(len(headers), len(values)))
                if headers[index]
            }
            parsed_rows.append(row)

        return parsed_rows

    raise ValueError(f"Unsupported file format: {suffix}. Use CSV, JSON, or PDF for market source.")


def _download_to_file(*, url: str, target_path: Path, user_agent: str) -> None:
    request = Request(url, headers={"User-Agent": user_agent})
    with urlopen(request, timeout=60.0) as response:
        data = response.read()
    target_path.write_bytes(data)


def _normalize_key(value: object) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def _normalize_scian_code(value: object) -> str | None:
    if value is None:
        return None

    text = str(value).strip()
    if not text:
        return None

    if re.fullmatch(r"\d{2}-\d{2}", text):
        return text

    if text.endswith(".0"):
        text = text[:-2]

    digits = re.sub(r"[^0-9]", "", text)
    return digits or None


def _infer_parent_code(
    code: str, merged_sectors: list[tuple[int, int, str]] | None = None
) -> str | None:
    if "-" in code or len(code) == 2:
        return None
    if len(code) == 3 and merged_sectors:
        prefix = int(code[:2])
        for low, high, merged_code in merged_sectors:
            if low <= prefix <= high:
                return merged_code
    return code[:-1]


def _clean_scian_title(value: str) -> str:
    """Strip whitespace and the trilateral 'T' marker INEGI appends to titles."""
    return re.sub(r"T\s*$", "", value.strip()).strip()


def _looks_like_scian_header_row(cells: list[str]) -> bool:
    normalized = {
        SCIAN_FIELD_ALIASES.get(_normalize_key(cell), "")
        for cell in cells
        if cell
    }
    return "code" in normalized and "title" in normalized
